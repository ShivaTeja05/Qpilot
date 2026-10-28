"""The wave 2 tools on a real lab file: the LightCycler 96 hmuY experiment of the corpus.

P. endodontalis hmuY and 16S rRNA, six samples in triplicate, SYBR Green I, melt curves in 0.2 C
steps (the depositors' own file, walked through Studio as a person would: the same walk as
``tests/web/test_melt_routes.py``). An agent reads its melt calls, each unit's value with its three
replicate wells, what to repeat and the numbers as a workbook; every number is compared with the
route the page reads.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from qpilot.mcp_server import QpilotTools
from tests.web import test_melt_routes as lc

lc96 = lc.lc96  # the walked file (a module-scoped fixture)
HMUY, PLATE, S16 = lc.HMUY, lc.PLATE, lc.S16


@pytest.fixture
def tools(lc96) -> QpilotTools:
    return QpilotTools(lc96["store"])


def refused(out: dict, code: str, field: str | None = None) -> dict:
    assert {"error", "code", "fix"} <= set(out), out
    assert out["code"] == code, out
    if field is not None:
        assert out.get("field") == field, out
    return out


def test_melt_calls_are_the_routes_calls_and_the_curves_stay_in_the_file(tools, lc96):
    c, study = lc96["client"], lc96["study"]
    route = lc.get(c, f"/api/studies/{study}/melt")
    out = tools.qpcr_melt(study)
    assert "error" not in out, out
    assert out["has_melt"] and out["plates"] == [PLATE] and out["pack"] == route["pack"]
    assert out["params"] == route["params"] and out["method"] == "methods/melt_curve.md"
    assert "MIQE 2.0" in out["caveat"] and out["saved"] is False
    assays = {a["target"]: a for a in out["assays"]}
    assert assays[HMUY]["counts"] == {"one_product": 18} and assays[HMUY]["flagged"] == []
    assert round(assays[HMUY]["expected_tm"], 1) == 82.0
    assert round(assays[S16]["expected_tm"], 1) == 87.1
    assert out["calls"] == route["calls"] and len(out["calls"]) == 96 and out["wells"] == []
    # no group is named in a melt reply
    assert not lc._names_a_group(out, set(lc.GROUPS.values()))
    # a well asked for gives its counted peaks, not the readings
    hmuy = sorted(x["well"] for x in out["calls"] if x["target"] == HMUY)[:2]
    one = tools.qpcr_melt(study, run=PLATE, wells=hmuy)
    assert sorted(w["well"] for w in one["wells"]) == sorted(hmuy)
    for w in one["wells"]:
        assert w["status"] == "one_product" and w["cq"] is not None
        (product,) = (p for p in w["peaks"] if p["role"] == "main")
        assert product["tm"] == w["main_tm"]
        assert not {"derivative", "temperature", "readings_temperature"} & set(w)
    concise = tools.qpcr_melt(study, response_format="concise")
    assert "calls" not in concise and concise["n_calls"] == 96 and concise["flagged"] == []
    # the session's id answers like its study's
    assert tools.qpcr_melt(lc96["sid"])["calls"] == out["calls"]


def test_melt_refusals(tools, lc96):
    study = lc96["study"]
    refused(tools.qpcr_melt(study, run="elsewhere"), "unknown_item", "run")
    refused(tools.qpcr_melt(study, wells=["A1", "../x"]), "invalid_argument", "wells")
    refused(tools.qpcr_melt("abcdef123456"), "unknown_study")


def test_a_session_is_blind_to_melt_until_both_stages_are_confirmed(tmp_path, monkeypatch):
    from qpilot.core.store import StudyStore

    monkeypatch.setenv("QPILOT_DATA_ROOTS", str(lc.LC96P.parent))
    monkeypatch.setenv("QPILOT_MODEL_PROVIDER", "none")
    tools = QpilotTools(StudyStore(tmp_path / "studies"))
    opened = tools.qpcr_inspect(str(lc.LC96P))
    assert "error" not in opened, opened
    out = refused(tools.qpcr_melt(opened["session_id"]), "results_hidden")
    assert "melt" not in out["error"].lower() or "confirm" in out["error"].lower()
    assert "peaks" not in str(out) and "Tm" not in str(out)


def test_units_show_the_three_replicate_wells_of_each_unit(tools, lc96):
    c, study = lc96["client"], lc96["study"]
    route = lc.get(c, f"/api/studies/{study}/genes/{HMUY}/units")
    out = tools.qpcr_units(study, HMUY)
    assert "error" not in out, out
    assert out["gene"] == HMUY and out["replicates"] is True and out["most_wells"] == 3
    assert [u["unit"] for u in out["units"]] == [p["unit"] for p in route["points"]]
    assert len(out["units"]) == 6 and {u["group"] for u in out["units"]} == {"Hm", "DIP"}
    for unit, point in zip(out["units"], route["points"], strict=True):
        assert unit["log2"] == point["x"] and unit["n_wells"] == point["n_wells"] == 3
        assert unit["cq_range"] == point["cq_range"]
        assert [w["well"] for w in unit["wells"]] == [w["well"] for w in point["wells"]]
        assert [w["cq"] for w in unit["wells"]] == [w["cq_value"] for w in point["wells"]]
        assert {w["state"] for w in unit["wells"]} == {"used"}
        # a unit's value is the mean of its wells' own values
        assert unit["log2"] == pytest.approx(sum(w["log2"] for w in unit["wells"]) / 3)
    assert out["fold_log2"] == route["fold_log2"] and out["p"] == route["p"]
    assert out["refs"] == [S16] and out["late_cq"] == 35.0
    first, second = out["groups"]
    assert {first["label"], second["label"]} == {"Hm", "DIP"}
    assert first["n"] == second["n"] == 3 and second["mean"] == pytest.approx(0.0)
    # the reference gene is not compared
    refused(tools.qpcr_units(study, S16), "not_applicable", "gene")


def test_the_repeat_plan_on_the_real_file_is_the_pages_plan(tools, lc96):
    from qpilot.web.routers import evidence as evidence_router

    evidence_router.clear()
    study = lc96["study"]
    out = tools.qpcr_repeat_plan(study)
    assert "error" not in out, out
    assert out["study_id"] == study and out["saved"] is False and out["include"] == "consequential"
    assert out["replicates"] >= 2 and out["plate_size"] in (96, 384)
    assert len(out["gate_lines"]) == 1 and out["gate_lines"][0]["plate"] == PLATE
    # every item names a reaction of this file and says why in words
    for item in out["items"]:
        assert item["sample"] in lc.GROUPS or item["scope"] == "sample"
        assert item["reasons"] and all("{{" not in r["words"] for r in item["reasons"])
    assert "person's call" in out["decision"]
    evidence_router.clear()


def test_the_workbook_and_figures_hold_the_units_and_the_wells(tools, lc96, tmp_path):
    import openpyxl

    study = lc96["study"]
    out = tools.qpcr_export(study, "xlsx")
    assert "error" not in out, out
    book = openpyxl.load_workbook(Path(out["file"]["path"]))
    assert book.sheetnames[:3] == ["Results", "Per unit", "Per well"]
    rows = list(book["Per unit"].iter_rows(values_only=True))
    header = list(rows[0])
    hmuy = [r for r in rows[1:] if r[header.index("Gene")] == HMUY]
    assert len(hmuy) == 6
    per_well = list(book["Per well"].iter_rows(values_only=True))
    assert len(per_well) - 1 >= 96
    figures = tools.qpcr_export(study, "figures", genes=[HMUY], basis="as_decided")
    assert "error" not in figures, figures
    assert figures["genes"] == [HMUY]
    import zipfile

    names = zipfile.ZipFile(io.BytesIO(Path(figures["file"]["path"]).read_bytes())).namelist()
    assert any(n.endswith(".svg") for n in names) and any(n.endswith(".png") for n in names)


def test_the_supplement_is_one_zip_with_the_record(tools, lc96):
    import zipfile

    out = tools.qpcr_export(lc96["study"], "supplement", genes=[HMUY])
    assert "error" not in out, out
    assert out["record"]["record_id"] and out["record"]["file"]["bytes"] > 0
    assert out["record"]["verify"].startswith("qpilot verify")
    names = zipfile.ZipFile(io.BytesIO(Path(out["file"]["path"]).read_bytes())).namelist()
    assert any(n.endswith(".xlsx") for n in names) and any(n.endswith("README.txt") for n in names)
