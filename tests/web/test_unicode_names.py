"""A lab's own group names, end to end through the page: every export keeps them, safely.

Real labs name their groups the way the bench does: ``LPS + inhibitor``, ``IL-6/ctrl``,
``Day 3 (50 µM)``, ``KO#2``, ``α-syn``. Before this fix any character outside letters, digits,
spaces and ``_ - . ( )`` made the Design page refuse the name, and a name set another way (an
agent, a sheet) made every export fail ("unsafe group name"): no workbook, CSV, figures,
supplement or Prism project, and the report itself failed for any study of three or more groups
(its Methods wrote the number of groups as a naked number).

Here a person sets exactly those five groups on the Design page (``POST …/design/set``) for a
mini file whose gene is ``IL-1β``, and the study goes through Curves to Results on the page's
routes. Then every export route is asked for its file: each is written, keeps the names as the
person wrote them (escaped in HTML and SVG, never a formula in a workbook or a CSV), and every
file name is plain ASCII (an ASCII slug and a short hash of the name where the name is not
already a plain file name), never a path.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from lxml import etree

from tests.web.conftest import BASE, TOKEN, wait_job

GROUPS = ("LPS + inhibitor", "IL-6/ctrl", "Day 3 (50 µM)", "KO#2", "α-syn")
CONTROL = "IL-6/ctrl"
GENE = "IL-1β"
SAMPLES = {f"S{i}": GROUPS[(i - 1) // 2] for i in range(1, 11)}  # two units per group
REASON = "groups from my lab notebook"
#: What any file name qPilot writes may be: ASCII letters, digits and ``_ - . ( )``.
SAFE_FILE = re.compile(r"^[A-Za-z0-9_(][A-Za-z0-9_().\-]*\Z")
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _plates() -> tuple[dict[str, str], dict[str, dict[str, list[float | None]]]]:
    samples = {s: "unkn" for s in SAMPLES} | {"NTC": "ntc"}
    ref = {s: [22.0 + 0.1 * (i % 2)] * 3 for i, s in enumerate(SAMPLES)}
    goi = {s: [25.0 + 0.7 * (i // 2) + 0.1 * (i % 2)] * 3 for i, s in enumerate(SAMPLES)}
    ref["NTC"], goi["NTC"] = [None] * 3, [None, None, 35.0]
    return samples, {"GAPDH": ref, GENE: goi}


def _items(*triples: tuple[str, str, dict]) -> list[dict]:
    return [{"kind": k, "subject": s, "payload": p} for k, s, p in triples]


def _head(c: TestClient, sid: str) -> str:
    r = c.get(f"/api/sessions/{sid}/design/editor")
    assert r.status_code == 200, r.text
    return r.json()["head"]


def _pending(c: TestClient, sid: str, stage: str) -> list[dict]:
    r = c.get(f"/api/sessions/{sid}/proposals", params={"stage": stage})
    return [p for p in r.json()["proposals"] if p["status"] == "proposed"]


def _confirm(c: TestClient, sid: str, stage: str) -> dict:
    view = c.post(f"/api/sessions/{sid}/confirm/{stage}").json()
    assert view["ready"], view
    r = c.post(
        f"/api/sessions/{sid}/confirm/{stage}/accept",
        json={"proposal_id": view["proposal"]["id"], "expected_head": view["head"]},
    )
    assert r.status_code == 200, r.text
    return r.json()


def _walk_on_the_page(c: TestClient, rdml: Path) -> tuple[str, str]:
    """Open, set the five groups and the comparisons on the Design page, Curves, Results."""
    r = c.post("/api/workspaces", files={"rdml": (rdml.name, rdml.read_bytes(), "x")})
    sid = wait_job(c, r.json()["job"])["result"]["key"]
    assert wait_job(c, c.post(f"/api/sessions/{sid}/propose/design", json={}).json()["job"])
    others = [g for g in GROUPS if g != CONTROL]
    comparison = {
        "group_a": others[0],
        "group_b": CONTROL,
        "pattern": "vs_control",
        "control": CONTROL,
        "correction": "holm",
        "comparisons": [{"group_a": g, "group_b": CONTROL} for g in others],
    }
    items = [
        *_items(("design.target_role", "GAPDH", {"role": "ref"})),
        *_items(("design.target_role", GENE, {"role": "toi"})),
        *_items(*(("design.group", s, {"group": g}) for s, g in SAMPLES.items())),
        *_items(("design.comparison", "comparison", comparison)),
    ]
    body = {"items": items, "reason": REASON, "expected_head": _head(c, sid)}
    r = c.post(f"/api/sessions/{sid}/design/set", json=body)
    assert r.status_code == 200, r.text
    for p in _pending(c, sid, "design"):
        r = c.post(
            f"/api/sessions/{sid}/proposals/{p['id']}/reject",
            json={"reason": "not part of this study", "expected_head": _head(c, sid)},
        )
        assert r.status_code == 200, r.text
    _confirm(c, sid, "design")
    assert wait_job(c, c.post(f"/api/sessions/{sid}/propose/curves").json()["job"])
    for p in _pending(c, sid, "curves"):
        r = c.post(
            f"/api/sessions/{sid}/proposals/{p['id']}/accept",
            json={"reason": "checked the curves", "expected_head": _head(c, sid)},
        )
        assert r.status_code == 200, r.text
    job = wait_job(c, _confirm(c, sid, "curves")["job"])
    assert job["status"] == "done", job
    return sid, job["result"]["study_id"]


@pytest.fixture(scope="module")
def named_walked(tmp_path_factory) -> dict[str, Any]:
    """The mini file with the five groups, walked to Results on the page's routes once."""
    from qpilot.ai.client import ModelClient, NoneBackend
    from qpilot.core.store import StudyStore
    from qpilot.qpcr import design_checks
    from qpilot.qpcr import inspect as qinspect
    from qpilot.studio_service import Studio
    from qpilot.web.context import make_context
    from qpilot.web.server import create_app
    from tests.fixtures.studio import studio_mini

    base = tmp_path_factory.mktemp("names")
    root = base / "data"
    root.mkdir()
    samples, plates = _plates()
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("QPILOT_DATA_ROOTS", str(root))
        mp.setenv("QPILOT_STUDIES_DIR", str(base / "studies"))
        mp.setattr(studio_mini, "SAMPLES", samples)
        mp.setattr(studio_mini, "PLATES", plates)
        mp.setattr(studio_mini, "TARGETS", {"GAPDH": "ref", GENE: "toi"})
        qinspect.clear_cache()
        rdml = studio_mini.write_studio_mini(root / "names.rdml")
        store = StudyStore(base / "studies")
        studio = Studio(store, client=ModelClient(backend=NoneBackend()))
        ctx = make_context(store, studio, token=TOKEN, watch=False)
        try:
            with TestClient(create_app(ctx), base_url=BASE) as c:
                c.headers["x-qpilot-token"] = TOKEN
                c.headers["origin"] = BASE
                sid, study = _walk_on_the_page(c, rdml)
        finally:
            ctx.close()
            design_checks._DECLARED.clear()
            design_checks.set_input_resolver(None)
            qinspect.clear_cache()
    return {"base": base, "root": root, "sid": sid, "study": study}


@pytest.fixture
def named(named_walked, tmp_path, monkeypatch) -> Iterator[dict[str, Any]]:
    import shutil

    from qpilot.ai.client import ModelClient, NoneBackend
    from qpilot.core.store import StudyStore
    from qpilot.qpcr import design_checks
    from qpilot.qpcr import inspect as qinspect
    from qpilot.studio_service import Studio
    from qpilot.web.context import make_context
    from qpilot.web.routers import evidence as evidence_router
    from qpilot.web.server import create_app

    shutil.copytree(named_walked["base"], tmp_path / "w", symlinks=False)
    monkeypatch.setenv("QPILOT_DATA_ROOTS", str(named_walked["root"]))
    monkeypatch.setenv("QPILOT_STUDIES_DIR", str(tmp_path / "w" / "studies"))
    qinspect.clear_cache()
    evidence_router.clear()
    store = StudyStore(tmp_path / "w" / "studies")
    studio = Studio(store, client=ModelClient(backend=NoneBackend()))
    ctx = make_context(store, studio, token=TOKEN, watch=False)
    with TestClient(create_app(ctx), base_url=BASE) as c:
        c.headers["x-qpilot-token"] = TOKEN
        c.headers["origin"] = BASE
        yield {"client": c, "store": store, "studio": studio, **named_walked}
    ctx.close()
    evidence_router.clear()
    design_checks._DECLARED.clear()
    design_checks.set_input_resolver(None)
    qinspect.clear_cache()


def _export(c: TestClient, study: str, fmt: str) -> tuple[dict, bytes]:
    r = c.post(f"/api/studies/{study}/report/export", json={"format": fmt})
    assert r.status_code == 202, r.text
    job = wait_job(c, r.json()["job"])
    assert job["status"] == "done", (fmt, job)
    link = job["result"]["download"]
    data = c.get(link["url"])
    assert data.status_code == 200, data.text
    return job["result"], data.content


def _live(text: str) -> list[str]:
    rows = csv.reader(io.StringIO(text.removeprefix("﻿")))
    cells = [v for row in rows for v in row]
    return [v for v in cells if v.startswith(FORMULA_START) and not _number(v)]


def _number(text: str) -> bool:
    try:
        float(text)
    except ValueError:
        return False
    return True


def _members(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _assert_safe_paths(names: list[str]) -> None:
    for name in names:
        parts = name.rstrip("/").split("/")
        assert all(SAFE_FILE.match(p) for p in parts), name


# ---- the design ---------------------------------------------------------------------------------


def test_the_design_page_keeps_the_names_as_written(named):
    c, sid = named["client"], named["sid"]
    view = c.get(f"/api/sessions/{sid}/design/editor").json()
    assert view["groups"] == {g: 2 for g in GROUPS}
    assert view["comparison"]["group_b"] == CONTROL


def test_a_name_nobody_can_read_is_still_refused_in_words(named):
    """Only an invisible or control character (or an empty or overlong name) is refused."""
    c, sid = named["client"], named["sid"]
    bad = ("KO‮2", "LPS\ninhibitor", "x" * 65)
    items = _items(*(("design.group", "S1", {"group": g}) for g in bad))
    plan = c.post(f"/api/sessions/{sid}/design/check", json={"items": items}).json()
    refusals = [p["refusal"] for p in plan["items"]]
    assert all(refusals), plan
    assert any("invisible character" in r for r in refusals)
    assert any("longer than 64" in r for r in refusals)
    assert "‮" not in str(refusals) and "\n" not in "".join(refusals)


# ---- the report and its tables ------------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["md", "html", "passport"])
def test_the_report_of_five_groups_is_written(named, fmt):
    result, data = _export(named["client"], named["study"], fmt)
    text = data.decode("utf-8")
    assert SAFE_FILE.match(result["download"]["name"]), result["download"]
    # every comparison is named in the Methods, as the person wrote the groups
    for group in GROUPS:
        assert group in text, (fmt, group)


def test_the_workbook_keeps_the_names_as_text(named):
    import openpyxl

    _, data = _export(named["client"], named["study"], "xlsx")
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    cells = [cell for s in book.worksheets for row in s.iter_rows() for cell in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    texts = {str(c.value) for c in cells if c.value is not None}
    assert set(GROUPS) <= texts, set(GROUPS) - texts
    assert GENE in texts


def test_the_csv_tables_keep_the_names_and_no_cell_is_live(named):
    _, data = _export(named["client"], named["study"], "csv")
    members = _members(data)
    _assert_safe_paths(list(members))
    joined = "".join(v.decode("utf-8") for v in members.values())
    assert all(g in joined for g in GROUPS) and GENE in joined
    for name, raw in members.items():
        if name.endswith(".csv"):
            assert not _live(raw.decode("utf-8")), name


def test_the_figures_carry_the_names_escaped_and_files_have_ascii_names(named):
    result, data = _export(named["client"], named["study"], "figures")
    assert result["genes"] == [GENE]
    members = _members(data)
    _assert_safe_paths(list(members))
    svgs = {n: v for n, v in members.items() if n.endswith(".svg")}
    assert svgs
    for name, raw in svgs.items():
        root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
        words = " ".join(t for t in root.itertext())
        if not name.endswith("panel.svg"):
            assert "LPS + inhibitor" in words and "IL-6/ctrl" in words, name
        tags = {etree.QName(e).localname for e in root.iter() if isinstance(e.tag, str)}
        assert not tags & {"script", "foreignObject", "a", "iframe", "image", "use"}, name


def test_the_supplement_is_written_with_ascii_file_names(named):
    import openpyxl

    _, data = _export(named["client"], named["study"], "supplement")
    members = _members(data)
    _assert_safe_paths(list(members))
    assert any(n.startswith("tables/per_gene/") for n in members)
    for name, raw in members.items():
        if name.endswith(".csv"):
            assert not _live(raw.decode("utf-8")), name
        if name.endswith(".xlsx"):
            book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
            assert not [
                c for s in book.worksheets for r in s.iter_rows() for c in r if c.data_type == "f"
            ]


def test_one_figure_is_drawn_and_its_file_name_is_ascii(named):
    c, study = named["client"], named["study"]
    info = c.get(f"/api/studies/{study}/report/figure/info", params={"gene": GENE})
    assert info.status_code == 200, info.text
    body = info.json()
    assert {g["name"] for g in body["groups"]} <= set(GROUPS)
    assert "α-syn vs IL-6/ctrl" in body["comparisons"]
    for fmt in ("svg", "png", "csv"):
        r = c.get(
            f"/api/studies/{study}/report/figure",
            params={"gene": GENE, "format": fmt, "comparison": "KO#2 vs IL-6/ctrl", "download": 1},
        )
        assert r.status_code == 200, (fmt, r.text[:200])
        name = r.headers["content-disposition"].split("filename=", 1)[1].strip('"')
        assert SAFE_FILE.match(name) and name.isascii(), name
        if fmt == "csv":
            assert "KO#2" in r.text and not _live(r.text)


# ---- the views ----------------------------------------------------------------------------------


def test_results_trace_and_units_show_the_names(named):
    c, study = named["client"], named["study"]
    comparisons = c.get(f"/api/studies/{study}/comparisons")
    assert comparisons.status_code == 200, comparisons.text
    found = comparisons.json()
    assert found["control"] == CONTROL
    labels = {(x["group_a"], x["group_b"]) for x in found["comparisons"]}
    assert labels == {(g, CONTROL) for g in GROUPS if g != CONTROL}
    assert "α-syn vs IL-6/ctrl" in {x["label"] for x in found["comparisons"]}
    assert len({x["id"] for x in found["comparisons"]}) == 4  # one id per comparison
    for route in ("results", f"genes/{GENE}", "trace", f"trace/{GENE}", f"genes/{GENE}/units"):
        r = c.get(f"/api/studies/{study}/{route}")
        if r.status_code == 409 and r.json()["error"]["code"] == "loading":
            wait_job(c, r.json()["error"]["details"]["job"])
            r = c.get(f"/api/studies/{study}/{route}")
        assert r.status_code == 200, (route, r.text[:300])
        assert "LPS + inhibitor" in r.text or "IL-6/ctrl" in r.text, route
    units = c.get(f"/api/studies/{study}/genes/{GENE}/units").json()
    shown = {units["groups"]["high"]["label"], units["groups"]["rest"]["label"]}
    assert shown <= set(GROUPS) and CONTROL in shown, shown


# ---- the files that go back to the bench --------------------------------------------------------


def test_the_repeat_plate_lays_out_the_names_as_they_are(named):
    c, study = named["client"], named["study"]
    for _ in range(200):
        r = c.get(f"/api/studies/{study}/review/repeat", params={"include": "all"})
        if r.status_code == 409 and r.json()["error"]["code"] == "loading":
            wait_job(c, r.json()["error"]["details"]["job"])
            continue
        break
    assert r.status_code == 200, r.text
    plan = r.json()["plan"]
    assert plan["wells"], plan["warnings"]
    assert not [w for w in plan["warnings"] if "cannot go into a plate layout" in w]
    text = c.get(f"/api/studies/{study}/review/repeat.csv", params={"include": "all"}).text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert {row["group"] for row in rows if row["group"]} <= set(GROUPS)
    assert {row["group"] for row in rows} & set(GROUPS)
    assert not _live(text)


def test_the_sample_sheet_template_round_trips_the_names(named):
    """The template carries the names; attached back as the sheet, the rules propose them."""
    c, sid = named["client"], named["sid"]
    r = c.get(f"/api/sessions/{sid}/design/template.csv")
    assert r.status_code == 200, r.text
    rows = list(csv.reader(io.StringIO(r.text)))
    written = {row[0]: row[1] for row in rows[1:] if row[0] in SAMPLES}
    assert written == SAMPLES
    from qpilot.qpcr.design_rules import _sheet_name

    assert {_sheet_name(g, "group") for g in GROUPS} == set(GROUPS)
    assert _sheet_name("'-RT", "group") == "-RT"  # the template's own defusing, read back


def test_the_prism_project_keeps_the_names(named):
    c, store, studio, sid = (named[k] for k in ("client", "store", "studio", "sid"))
    study = named["study"]
    for card in store.cards(study):
        if card.status in ("open", "pending_human"):
            head = store.chain_head(study)
            studio.decide(sid, card.id, card.recommended, "the planned option", "human", head)
    studio.accept(sid, studio.propose_confirm(sid, "choices").proposal.id, "human")
    preview = c.get(f"/api/sessions/{sid}/prism/preview")
    assert preview.status_code == 200, preview.text
    drafted = wait_job(c, c.post(f"/api/sessions/{sid}/prism/draft").json()["job"])
    assert drafted["status"] == "done", drafted
    studio.accept(sid, drafted["result"]["proposal_id"], "human")
    sent = wait_job(c, c.post(f"/api/sessions/{sid}/prism/send").json()["job"], 600)
    assert sent["status"] == "done", sent
    folder = Path(sent["result"]["bundle"]["folder"])
    _assert_safe_paths(list(sent["result"]["files"]))
    from qpilot.prism import pzfx

    project = pzfx.read_project((folder / sent["result"]["project"]).read_bytes())
    titles = {name for t in project.tables for name in t.column_names()}
    assert set(GROUPS) <= titles, titles
