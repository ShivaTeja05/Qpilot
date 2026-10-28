"""The Excel workbook as a Report export (origin/main's ``qpilot export excel``, on the page)."""

from __future__ import annotations

import io
import zipfile

from tests.web.conftest import wait_job


def test_the_workbook_is_a_report_export_with_the_results_and_where_to_look(walked):
    import openpyxl

    c, store, study = walked["client"], walked["store"], walked["study"]
    before = store.chain_head(study)
    r = c.post(
        f"/api/studies/{study}/report/export",
        json={"format": "xlsx", "expected_head": before},
    )
    assert r.status_code == 202, r.text
    job = wait_job(c, r.json()["job"])
    assert job["status"] == "done", job
    out = job["result"]
    assert out["format"] == "xlsx" and out["download"]["name"].endswith(".xlsx")
    assert out["head_moved"] is False and store.chain_head(study) == before
    data = c.get(out["download"]["url"])
    assert data.status_code == 200
    zipfile.ZipFile(io.BytesIO(data.content)).testzip()
    book = openpyxl.load_workbook(io.BytesIO(data.content))
    assert book.sheetnames[:3] == ["Results", "Per unit", "Per well"]
    assert "Where to look" in book.sheetnames and book.sheetnames[-1] == "About"
    rows = list(book["Results"].iter_rows(values_only=True))
    assert rows[0][-1] == "Trace" and len(rows) - 1 == len(store.get(study).contrasts)
    about = dict(book["About"].iter_rows(values_only=True))
    assert about["study_id"] == store.get(study).study_id
    assert out["bases"] == ["as_decided"]
    listed = c.get(f"/api/studies/{study}/report/exports").json()["exports"]
    assert listed[0]["format"] == "xlsx"


def test_an_unknown_format_is_refused(walked):
    c, study = walked["client"], walked["study"]
    assert c.post(f"/api/studies/{study}/report/export", json={"format": "xls"}).status_code == 422
