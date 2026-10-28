"""Excel export of a stored study: its numbers as numbers, one sheet per table.

:func:`export_excel` writes one ``.xlsx`` workbook from :func:`qpilot.report.tables.study_tables`:

* ``Results``: one row per comparison and gene, as decided on the Decision Cards: the call, the
  fold change and its 95 % CI as separate numeric columns, log2 fold change, p, q (when the
  engine gives one), n per group, the share of qPilot's choice grid with the same call, the
  confidence level, and a ``Trace`` column naming each gene's top place to look and what it
  does (e.g. ``HMBS_20080318/K22 (2016): changes the call``, or ``nothing unusual``);
* ``Per unit``: ΔCq, log2 RQ and RQ per biological unit (the values the test uses);
* ``Per well``: plate, well, sample, target, Cq, N0, efficiency, the amplification call, and
  whether a person left the well out and why;
* ``Results (recommended)`` and ``Per unit (recommended)`` when the options chosen are not the
  recommended ones (both analyses, each labelled);
* ``Where to look``: one row per place Trace ranks for each gene (the ``top`` first places);
* ``Decisions`` (who, what, why, when), ``Settings`` (every choice in use, the test, the rule),
  ``Methods`` (the report's paragraphs) and ``About`` (the study, the analysis, the input's and
  the design's sha256, the decision chain head the Analysis Record carries).

Every number is a number cell with a number format (fold changes ``0.000``, p ``0.00E+00``
below 0.001, shares as percentages, Cq ``0.00``), headers name the units, the header row is
frozen (and the columns that name a row), columns have widths and a filter.

Trace runs read-only from the study's input file (as :func:`qpilot.qpcr.trace.trace_store`
does); when it cannot (the input file is unknown or changed), the results are still written,
the ``Trace`` column says ``not traced`` and the ``About`` sheet says why. Nothing is written to
the study's workspace, and the export decides nothing.

The file is byte-deterministic: openpyxl's own save stamps the time, so the package is
rewritten with fixed ``docProps/core.xml`` dates, members in sorted order and fixed member dates.
Text that a spreadsheet would read as a formula (``=``, ``+``, ``-``, ``@`` first) is stored as
text. Deterministic: no model call and no network.
"""

from __future__ import annotations

import io
import os
import re
import tempfile
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

RESULTS_SHEET = "Results"
WHERE_SHEET = "Where to look"
ABOUT_SHEET = "About"
TRACE_COLUMN = "Trace"
WHERE_COLUMNS = (
    "result",
    "rank",
    "sample",
    "plate",
    "well",
    "checks",
    "consequence",
    "fold_off",
    "changes_call",
    "z",
    "calls_changed",
)
NOT_TRACED = "not traced"
RESEARCH_USE = "For research use only. Not for clinical or diagnostic use."
FIXED_TIME = b"2000-01-01T00:00:00Z"  # docProps/core.xml created and modified
ZIP_DATE = (1980, 1, 1, 0, 0, 0)  # the earliest date a zip member can carry
FORMULA_START = ("=", "+", "-", "@")
DIGITS = 6  # decimals kept for fold_off and z in the "Where to look" sheet
_CORE_DATES = re.compile(rb"<dcterms:(created|modified)([^>]*)>[^<]*<")

#: Number formats by column kind (:class:`qpilot.report.tables.Column`).
NUMBER_FORMATS: dict[str, str] = {
    "int": "0",
    "float": "0.0000",
    "fold": "0.000",
    "log2": "0.000",
    "rq": "0.0000",
    "p": "[<0.001]0.00E+00;0.0000",
    "share": "0.0%",
    "cq": "0.00",
    "n0": "0.000E+00",
    "eff": "0.000",
}
#: Columns kept in view while scrolling right: (sheet name, first scrolling cell).
FREEZE: dict[str, str] = {
    "results": "E2",
    "results_recommended": "E2",
    "per_unit": "E2",
    "per_unit_recommended": "E2",
    "per_well": "C2",
}
WIDE = {"value": 90, "text": 120, "why": 60, "evidence": 80, "note": 44, "left_out_reason": 44}
HEADER_FILL = "FFE8EEF4"


class ExcelError(ValueError):
    """The Excel file cannot be written: a bad path, an existing file, or openpyxl missing."""


@dataclass(frozen=True)
class ExcelExport:
    """What was written: the file, its row counts, its sheets and how Trace ran."""

    path: Path
    n_results: int
    n_places: int
    traced: bool
    note: str = ""
    sheets: tuple[str, ...] = ()
    bases: tuple[str, ...] = ()


def _openpyxl():
    try:
        import openpyxl
    except ImportError as err:  # an optional extra, like reading .xlsx sheets
        raise ExcelError("writing .xlsx needs openpyxl (pip install 'qpilot[studio]')") from err
    return openpyxl


def _traces(store, study_id: str, genes: Sequence[str], top: int, context) -> tuple[dict, str]:
    """Each gene's trace (``top`` places), or ``{}`` and why Trace could not run."""
    from qpilot.qpcr.trace import TraceError, _store_context, trace_gene

    if context is None:
        try:
            context = _store_context(store, study_id)
        except TraceError as err:
            return {}, str(err)
        except Exception as err:  # the results are still exported; the About sheet says why
            return {}, f"Trace could not run ({type(err).__name__})"
    out = {}
    for gene in genes:
        try:
            out[gene] = trace_gene(context, gene, top=top)
        except TraceError:
            continue  # a gene Trace cannot follow: its cell says "not traced"
    return out, ""


def _num(value: Any) -> float | None:
    if value is None:
        return None
    number = round(float(value), DIGITS)
    return number if number == number else None  # NaN stays empty


def where_rows(traces: Mapping[str, Any]) -> list[tuple]:
    """The "Where to look" rows: every ranked place of every gene, in the order given."""
    from qpilot.explain.trace_text import CONSEQUENCE_WORDS

    rows = []
    for gene, result in traces.items():
        for place in result.places:
            option = place.options.get(place.via, {})
            rows.append(
                (
                    gene,
                    place.rank,
                    place.sample,
                    place.plate,
                    place.well,
                    " ".join(place.checks) or None,
                    CONSEQUENCE_WORDS.get(place.consequence, place.consequence),
                    _num(place.fold_off),
                    place.consequence == "changes_call",
                    _num(place.z),
                    ", ".join(option.get("changed_calls") or ()) or None,
                )
            )
    return rows


def _clean(value: Any) -> Any:
    """A cell value: text without characters XML forbids (empty text is no value)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value) or None
    return value


def _cell(sheet, value: Any, fmt: str | None = None, bold: bool = False):
    """One write-only cell; text never becomes a formula."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    value = _clean(value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str) and value.startswith(FORMULA_START):
        cell.data_type = "s"  # text from a file is never a formula
    if fmt and isinstance(value, int | float) and not isinstance(value, bool):
        cell.number_format = fmt
    if bold:
        cell.font = Font(bold=True)
    return cell


def _header_cells(sheet, headers: Sequence[str]) -> list:
    from openpyxl.styles import Alignment, PatternFill

    out = []
    for text in headers:
        cell = _cell(sheet, text, bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        out.append(cell)
    return out


def _width(header: str, kind: str, key: str) -> float:
    if key in WIDE:
        return WIDE[key]
    if kind == "text":
        return min(max(len(header) * 0.9, 12), 40)
    return min(max(len(header) * 0.55, 10), 22)


def _put_table(book, table, extra: Sequence[tuple[str, Sequence[Any]]] = ()) -> Any:
    """One table as a sheet: a bold, frozen, filtered header row and typed number cells.

    ``extra`` are more text columns, each ``(header, values)``, appended on the right.
    """
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(table.title[:31])
    columns = list(table.columns)
    headers = [c.header for c in columns] + [h for h, _ in extra]
    kinds = [c.kind for c in columns] + ["text"] * len(extra)
    keys = [c.key for c in columns] + [h.lower() for h, _ in extra]
    for i, (header, kind, key) in enumerate(zip(headers, kinds, keys, strict=True), start=1):
        sheet.column_dimensions[get_column_letter(i)].width = _width(header, kind, key)
    sheet.freeze_panes = FREEZE.get(table.name, "A2")
    n_rows = len(table.rows)
    last = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last}{max(n_rows, 1) + 1}"
    sheet.append(_header_cells(sheet, headers))
    formats = [NUMBER_FORMATS.get(k) for k in kinds]
    extra_values = [list(values) for _, values in extra]
    for r, row in enumerate(table.rows):
        values = [*row, *(col[r] for col in extra_values)]
        sheet.append([_cell(sheet, v, f) for v, f in zip(values, formats, strict=True)])
    if not n_rows and table.note:
        sheet.append([_cell(sheet, f"No rows: {table.note}")])
    return sheet


def _put_rows(sheet, rows: Iterable[Sequence[Any]], header: bool = True) -> None:
    """Write plain rows (the "Where to look" sheet): the first row bold when ``header``."""
    for r, row in enumerate(rows):
        if r == 0 and header:
            sheet.append(_header_cells(sheet, [str(v) for v in row]))
        else:
            sheet.append([_cell(sheet, v) for v in row])


def deterministic_xlsx(book) -> bytes:
    """``book`` saved as bytes that depend only on its content (fixed dates, sorted members)."""
    raw, out = io.BytesIO(), io.BytesIO()
    book.save(raw)
    with zipfile.ZipFile(raw) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in sorted(zin.infolist(), key=lambda i: i.filename):
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _CORE_DATES.sub(rb"<dcterms:\1\2>" + FIXED_TIME + b"<", data)
            member = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE)
            member.compress_type = zipfile.ZIP_DEFLATED
            member.external_attr = 0o600 << 16
            zout.writestr(member, data)
    return out.getvalue()


def excel_bytes(
    store,
    study_id: str,
    *,
    top: int = 3,
    context=None,
    tables=None,
    record_id: str | None = None,
    trace: bool = True,
) -> tuple[bytes, ExcelExport]:
    """The workbook's bytes and what it holds (``ExcelExport.path`` is empty).

    Args:
        store: The study store.
        study_id: The study.
        top: Places per gene in the "Where to look" sheet (1 to 50).
        context: A :class:`qpilot.qpcr.trace.TraceContext` already built for this study
            (default: built read-only from the study's input file).
        tables: The study's :class:`~qpilot.report.tables.StudyTables` (default: read here).
        record_id: The Analysis Record written with this workbook, named on the About sheet.
        trace: Run Trace for the ``Trace`` column and the "Where to look" sheet.

    Raises:
        ExcelError: ``top`` out of range, or openpyxl is not installed.
        qpilot.core.store.StudyNotFound: Unknown study.
    """
    from qpilot.app_state import trace_cell
    from qpilot.report.tables import study_tables

    if isinstance(top, bool) or not isinstance(top, int) or not 1 <= top <= 50:
        raise ExcelError("top must be a whole number from 1 to 50")
    openpyxl = _openpyxl()
    tables = tables if tables is not None else study_tables(store, study_id)
    decided = tables.decided
    genes = list(dict.fromkeys(r["gene"] for r in decided.results.records()))
    if trace:
        traces, note = _traces(store, study_id, genes, top, context)
    else:
        traces, note = {}, "not asked for"
    gene_at = decided.results.keys().index("gene")
    trace_column = [
        trace_cell(traces[row[gene_at]]) if row[gene_at] in traces else NOT_TRACED
        for row in decided.results.rows
    ]
    book = openpyxl.Workbook(write_only=True)
    book.properties.creator = "qPilot"
    _put_table(book, decided.results, extra=[(TRACE_COLUMN, trace_column)])
    _put_table(book, decided.per_unit)
    _put_table(book, decided.per_well)
    if tables.recommended is not None:
        _put_table(book, tables.recommended.results)
        _put_table(book, tables.recommended.per_unit)
    where = book.create_sheet(WHERE_SHEET)
    where.freeze_panes = "A2"
    places = where_rows(traces)
    _put_rows(where, [WHERE_COLUMNS, *places])
    for table in (tables.decisions, tables.settings, tables.methods):
        _put_table(book, table)
    how = (
        f"Trace ran on {len(traces)} of {len(genes)} results, at most {top} places each"
        if traces
        else f"Trace did not run: {note or 'no result could be traced'}"
    )
    about = tables.about(record_id)
    rows = [*about.rows]
    rows.insert(-1, ("trace", how))
    _put_table(book, replace(about, title=ABOUT_SHEET, rows=tuple(rows)))
    data = deterministic_xlsx(book)
    sheets = [decided.results.title, decided.per_unit.title, decided.per_well.title]
    if tables.recommended is not None:
        sheets += [tables.recommended.results.title, tables.recommended.per_unit.title]
    sheets += [WHERE_SHEET, tables.decisions.title, tables.settings.title, tables.methods.title]
    sheets.append(ABOUT_SHEET)
    info = ExcelExport(
        Path(),
        len(genes),
        len(places),
        bool(traces),
        note,
        sheets=tuple(sheets),
        bases=tables.bases,
    )
    return data, info


def write_atomic(out: Path, data: bytes) -> None:
    """Write ``data`` to ``out`` through a temporary file in the same folder (all or nothing)."""
    fd, tmp = tempfile.mkstemp(dir=out.parent, prefix=f".{out.name}.", suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp, out)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def new_file(out_path: str | Path, suffix: str, what: str = "file") -> Path:
    """A path for a new file: the right suffix, not there yet, in an existing folder.

    Raises:
        ExcelError: Wrong suffix, the file exists, or its folder does not.
    """
    out = Path(out_path).expanduser()
    if out.suffix.lower() != suffix:
        raise ExcelError(f"the {what} must end in {suffix}")
    if out.exists():
        raise ExcelError(f"{out} already exists; choose a new file")
    if not out.parent.is_dir():
        raise ExcelError(f"the folder {out.parent} does not exist")
    return out


def export_excel(
    store,
    study_id: str,
    out_path: str | Path,
    *,
    top: int = 3,
    context=None,
    tables=None,
    record_id: str | None = None,
) -> ExcelExport:
    """Write the study's tables and Trace's places to a new ``.xlsx`` file.

    Args:
        store: The study store (read only).
        study_id: The study.
        out_path: A new file ending in ``.xlsx`` in an existing folder.
        top: Places per gene in the "Where to look" sheet.
        context: A Trace context already built for this study (optional).
        tables: The study's tables (default: read here).
        record_id: The Analysis Record written with it (named on the About sheet).

    Returns:
        What was written.

    Raises:
        ExcelError: The path does not end in ``.xlsx``, the file exists, its folder does not,
            ``top`` is out of range, or openpyxl is not installed.
        qpilot.core.store.StudyNotFound: Unknown study.
    """
    out = new_file(out_path, ".xlsx", "Excel file")
    data, info = excel_bytes(
        store, study_id, top=top, context=context, tables=tables, record_id=record_id
    )
    write_atomic(out, data)
    return ExcelExport(
        out, info.n_results, info.n_places, info.traced, info.note, info.sheets, info.bases
    )
