"""What the e2e specs check inside the files the page downloads (run with the venv's python3 -I).

    files.py xlsx-numbers BOOK.xlsx GENE     the Results sheet has GENE's row; its fold change,
                                             CI, p and n cells are number cells (prints JSON)
    files.py zip-names ARCHIVE.zip           the members of a zip (prints JSON)
    files.py pzfx-exclude BUNDLE.zip OUT     the bundle's Prism project with the first value of
                                             its first table excluded, as Prism saves it (prints
                                             the table title and the row)

Standard library and openpyxl/lxml only; nothing here imports qPilot.
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

NS = "http://graphpad.com/prism/Prism.htm"


def xlsx_numbers(book: str, gene: str) -> dict:
    """The sheet names, and the type and value of each numeric column of GENE's Results row."""
    import openpyxl

    wb = openpyxl.load_workbook(book, read_only=False)
    sheet = wb["Results"]
    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = [r for r in sheet.iter_rows(min_row=2) if r[header.index("Gene")].value == gene]
    if not rows:
        raise SystemExit(f"no row for {gene} in Results: {header}")
    cell = dict(zip(header, rows[0], strict=False))
    numeric = {}
    for key, c in cell.items():
        if not isinstance(key, str):
            continue
        if key.startswith(("Fold change", "95 % CI", "p (", "n ")):
            numeric[key] = {"type": c.data_type, "value": c.value}
    return {"sheets": wb.sheetnames, "gene_type": cell["Gene"].data_type, "numeric": numeric}


def zip_names(archive: str) -> list[str]:
    """The member names of a zip archive, sorted."""
    with zipfile.ZipFile(archive) as z:
        return sorted(z.namelist())


def pzfx_exclude(bundle: str, out: str) -> dict:
    """Write the bundle's project to OUT with one value excluded, as Prism marks it."""
    from lxml import etree

    with zipfile.ZipFile(bundle) as z:
        name = next(n for n in z.namelist() if n.endswith(".pzfx"))
        root = etree.fromstring(z.read(name))
    ns = {"p": NS}
    table = root.find("p:Table", ns)
    row = table.find("p:RowTitlesColumn/p:Subcolumn/p:d", ns).text
    for column in table.findall("p:YColumn", ns):
        cell = column.find("p:Subcolumn/p:d", ns)
        if cell is not None and cell.text:
            cell.set("Excluded", "1")
            break
    Path(out).write_bytes(etree.tostring(root, xml_declaration=True, encoding="UTF-8"))
    return {"project": name, "table": table.find("p:Title", ns).text, "row": row}


def main(argv: list[str]) -> None:
    """Run one command and print its answer as JSON."""
    command, *args = argv
    done = {"xlsx-numbers": xlsx_numbers, "zip-names": zip_names, "pzfx-exclude": pzfx_exclude}
    print(json.dumps(done[command](*args)))


if __name__ == "__main__":
    main(sys.argv[1:])
